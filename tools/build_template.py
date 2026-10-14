"""Build the tagged EHS-003a template used by the generator.

Reads EHS's original form and writes a copy with docxtpl (Jinja) tags in
place of the instruction text. Re-run this whenever EHS issues a new revision
of the form, then check the result in Word before using it.

    python tools/build_template.py "EHS-003a LOTO Machine-Specific Energy.docx" templates/EHS-003a_v1.docx

Add --keep-original-widths to leave the Description column widths exactly as on the form.

Structural changes made to the form:
  * Description label/value columns are rebalanced so values have room
    (see WIDE_GRID; disable with --keep-original-widths).
  * Description values, step boxes, restart box and footer get tags.
  * Instruction text, the "Example:" rows and the example point are removed.
  * The single big table is split into three tables (top / isolation points /
    bottom) so the isolation-point header row repeats on every page and a
    point row never splits across a page break.
"""
import copy
import sys
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W_T = qn("w:t")


def cell_text(tc):
    return "".join(t.text or "" for t in tc.iter(W_T)).strip()


def row_cells(tr):
    return tr.findall(qn("w:tc"))


def set_cell(tc, text, italic=False, keep_paragraphs=1):
    """Replace a cell's content with a single run of `text`.

    Keeps the first paragraph's formatting and the first run's character
    formatting (minus italics) so tags inherit the form's look.
    """
    paras = tc.findall(qn("w:p"))
    p = paras[0]
    for extra in paras[keep_paragraphs:]:
        tc.remove(extra)
    runs = p.findall(qn("w:r"))
    rpr = None
    if runs and runs[0].find(qn("w:rPr")) is not None:
        rpr = copy.deepcopy(runs[0].find(qn("w:rPr")))
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    r = OxmlElement("w:r")
    if rpr is None:
        ppr_rpr = p.find(qn("w:pPr") + "/" + qn("w:rPr"))
        rpr = copy.deepcopy(ppr_rpr) if ppr_rpr is not None else OxmlElement("w:rPr")
    if not italic:
        for tag in ("w:i", "w:iCs"):
            for el in rpr.findall(qn(tag)):
                rpr.remove(el)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    # Drop any list numbering on the value paragraph.
    ppr = p.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        ppr.remove(ppr.find(qn("w:numPr")))
    return p


def add_paragraph_after(p, text):
    new = copy.deepcopy(p)
    for child in list(new):
        if child.tag != qn("w:pPr"):
            new.remove(child)
    r = OxmlElement("w:r")
    rpr = p.find(qn("w:r") + "/" + qn("w:rPr"))
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    new.append(r)
    p.addnext(new)
    return new


def set_row_prop(tr, tag):
    trpr = tr.find(qn("w:trPr"))
    if trpr is None:
        trpr = OxmlElement("w:trPr")
        tr.insert(0, trpr)
    if trpr.find(qn(tag)) is None:
        trpr.append(OxmlElement(tag))


def spacer_paragraph():
    """Near-zero-height paragraph; Word needs one between adjacent tables."""
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "20")
    spacing.set(qn("w:lineRule"), "exact")
    ppr.append(spacing)
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    rpr.append(sz)
    ppr.append(rpr)
    p.append(ppr)
    return p


def find_row(rows, startswith):
    for i, tr in enumerate(rows):
        if cell_text(tr).startswith(startswith):
            return i
    raise ValueError(f"Row starting with {startswith!r} not found - has the form changed?")


# Description block column widths (twips). The original value boxes are ~0.9"
# and ~1.3" wide, which squeezes the energy summary into a tall column. These
# widths keep the table width and the 50/50-ish split of the numbered step
# rows (cols 0-2 | cols 3-5) unchanged; only the label/value split moves.
ORIGINAL_GRID = [3189, 329, 933, 1808, 1325, 1856]
WIDE_GRID = [2200, 1318, 933, 900, 1325, 2764]


def widen_description(tbl):
    grid = [int(g.get(qn("w:w"))) for g in tbl.find(qn("w:tblGrid"))]
    if grid != ORIGINAL_GRID:
        raise ValueError(f"Unexpected column grid {grid} - has the form changed?")
    for g, w in zip(tbl.find(qn("w:tblGrid")), WIDE_GRID):
        g.set(qn("w:w"), str(w))
    for tr in tbl.findall(qn("w:tr")):
        col = 0
        for tc in row_cells(tr):
            span_el = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
            span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            tcw = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
            if tcw is not None:
                tcw.set(qn("w:w"), str(sum(WIDE_GRID[col:col + span])))
                tcw.set(qn("w:type"), "dxa")
            col += span


def tag_body(document, widen=True):
    body = document.element.body
    tbl = body.find(qn("w:tbl"))
    rows = tbl.findall(qn("w:tr"))

    # Description block: label cell followed by value cell.
    values = {
        "Facility/Site:": "{{ site_name }}",
        "Department:": "{{ department_name }}",
        "Machine/Equipment:": "{{ machine_equipment }}",
        "Tasks Covered:": "{{r tasks_covered }}",
        "Energy Types & Magnitudes Involved:": "{{r energy_summary }}",
        "# of Locks Required:": "{{ locks_required }}",
    }
    for tr in rows[:6]:
        cells = row_cells(tr)
        for i, tc in enumerate(cells[:-1]):
            label = cell_text(tc)
            if label in values:
                set_cell(cells[i + 1], values.pop(label))
    if values:
        raise ValueError(f"Description labels not found: {list(values)}")

    # Numbered step boxes: label on the left, value box on the right.
    steps = {
        "Notify Affected Employees": "{{r affected_employees }}",
        "Shutdown Equipment Using Documented Procedures": "{{r shutdown_procedure }}",
        "Isolate All Energy Sources": "{{r isolate_text }}",
        "Apply Lockout/Tagout Devices": "{{r lockout_devices }}",
        "Release or Restrain Stored Energy": "{{r stored_energy }}",
        "Verify Isolation Prior to Beginning Work": "{{r verify_isolation }}",
        "Perform Servicing or Maintenance": "{{r servicing_notes }}",
    }
    for tr in rows:
        cells = row_cells(tr)
        if len(cells) == 2 and cell_text(cells[0]) in steps:
            set_cell(cells[1], steps.pop(cell_text(cells[0])))
    if steps:
        raise ValueError(f"Step rows not found: {list(steps)}")

    # Restart steps go in the empty box under "Release from Lockout".
    i_release = find_row(rows, "Release from Lockout")
    set_cell(row_cells(rows[i_release + 1])[0], "{{r restart_steps }}")

    # Isolation table: header, (instructions), (Example:), header, example row, blank.
    i_header = find_row(rows, "Picture of Energy")
    i_note = find_row(rows, "Energy types and isolation points must be listed")
    iso_rows = rows[i_header:i_note]
    header, instructions, example_label, for_row, point_row, endfor_row = iso_rows
    tbl.remove(instructions)
    tbl.remove(example_label)

    # for / endfor rows: docxtpl removes these rows when rendering.
    set_cell(row_cells(for_row)[0], "{%tr for p in points %}")
    for tc in row_cells(for_row)[1:]:
        set_cell(tc, "")
    set_cell(row_cells(endfor_row)[0], "{%tr endfor %}")
    for tc in row_cells(endfor_row)[1:]:
        set_cell(tc, "")

    photo_tc, energy_tc, control_tc = row_cells(point_row)
    # Photo cell: one paragraph per photo (or a "photo pending" note).
    p = set_cell(photo_tc, "{%p for ph in p.photos %}")
    p2 = add_paragraph_after(p, "{{ ph.image }}")
    p3 = add_paragraph_after(p2, "{%p endfor %}")
    add_paragraph_after(p3, "{{r p.photo_note }}")
    # Energy cell: point label + "Type – magnitude", color applied after render.
    p = set_cell(energy_tc, "{{r p.energy_label }}")
    # Control information: one bulleted paragraph per line, list style kept.
    first = control_tc.findall(qn("w:p"))[0]
    for extra in control_tc.findall(qn("w:p"))[1:]:
        control_tc.remove(extra)
    for child in list(first):
        if child.tag != qn("w:pPr"):
            first.remove(child)
    lp = copy.deepcopy(first)
    numpr = lp.find(qn("w:pPr") + "/" + qn("w:numPr"))
    if numpr is not None:
        lp.find(qn("w:pPr")).remove(numpr)
    loop_start = copy.deepcopy(lp)
    loop_end = copy.deepcopy(lp)
    for el, text in ((loop_start, "{%p for line in p.control_lines %}"), (first, "{{r line }}"), (loop_end, "{%p endfor %}")):
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        el.append(r)
    first.addprevious(loop_start)
    first.addnext(loop_end)

    set_row_prop(header, "w:tblHeader")
    set_row_prop(point_row, "w:cantSplit")

    # Split into three tables so the header row can repeat on each page.
    def new_table_like(src):
        t = OxmlElement("w:tbl")
        t.append(copy.deepcopy(src.find(qn("w:tblPr"))))
        t.append(copy.deepcopy(src.find(qn("w:tblGrid"))))
        return t

    iso_tbl = new_table_like(tbl)
    bottom_tbl = new_table_like(tbl)
    for tr in (header, for_row, point_row, endfor_row):
        iso_tbl.append(tr)
    for tr in rows[i_note:]:
        bottom_tbl.append(tr)
    tbl.addnext(spacer_paragraph())
    tbl.getnext().addnext(iso_tbl)
    iso_tbl.addnext(spacer_paragraph())
    iso_tbl.getnext().addnext(bottom_tbl)

    if widen:
        widen_description(tbl)  # top table only; isolation/bottom tables keep the original grid


def tag_footer(document):
    footer = document.sections[0].footer._element
    replacements = {
        "Document Effective Date": "{{ effective_date }}",
        "Document Review Date": "{{ review_date }}",
        "Document Owner": "{{ document_owner }}",
    }
    found = 0
    for tc in footer.iter(qn("w:tc")):
        label = cell_text(tc)
        for key, tag in replacements.items():
            if label.startswith(key):
                # Label on the first line, value on the second (as on the form).
                label_p = set_cell(tc, f"{key}:", italic=True)
                add_paragraph_after(label_p, tag)
                found += 1
    if found != len(replacements):
        raise ValueError("Footer date/owner cells not found - has the form changed?")
    # Procedure number / status line in the empty paragraph above the table.
    for p in footer.findall(qn("w:p")):
        if not cell_text(p):
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), "18")
            rpr.append(sz)
            r.append(rpr)
            t = OxmlElement("w:t")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t.text = "Procedure No. {{ procedure_number }}   |   Status: {{ status_label }}"
            r.append(t)
            p.append(r)
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), "center")
            p.find(qn("w:pPr")).append(jc)
            break


def main(src, dst, widen=True):
    document = docx.Document(src)
    tag_body(document, widen=widen)
    tag_footer(document)
    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    document.save(dst)
    print(f"Wrote {dst}")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 2:
        sys.exit(__doc__)
    main(args[0], args[1], widen="--keep-original-widths" not in sys.argv)
